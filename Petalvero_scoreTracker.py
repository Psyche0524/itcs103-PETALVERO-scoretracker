import openpyxl
from openpyxl import Workbook, load_workbook
import os

FILENAME = "student_scores.xlsx"

def get_status(score):
    return "Pass" if score >= 50 else "Fail"

def save_data(name, score):
    if os.path.exists(FILENAME):
        wb = load_workbook(FILENAME)
        sheet = wb.active
    else:
        wb = Workbook()
        sheet = wb.active
        sheet.append(["Name", "Score", "Status"])  # Header

    status = get_status(score)
    sheet.append([name, score, status])
    wb.save(FILENAME)
    print("Data saved successfully.\n")

def show_all_records():
    if not os.path.exists(FILENAME):
        print("No records found.")
        return

    wb = load_workbook(FILENAME)
    sheet = wb.active
    print("\nAll Records:")
    for row in sheet.iter_rows(min_row=2, values_only=True):
        print(f"Name: {row[0]}, Score: {row[1]}, Status: {row[2]}")
    print()

def main():
    while True:
        print("1. Add Student Score")
        print("2. View All Records")
        print("3. Exit")
        choice = input("Enter your choice: ")

        if choice == "1":
            name = input("Enter student name: ")
            try:
                score = float(input("Enter score: "))
                save_data(name, score)
            except ValueError:
                print("Invalid score. Please enter a number.\n")
        elif choice == "2":
            show_all_records()
        elif choice == "3":
            print("Exiting program.")
            break
        else:
            print("Invalid choice. Try again.\n")

if __name__ == "__main__":
    main()
